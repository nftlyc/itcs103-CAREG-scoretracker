import tkinter as tk
from tkinter import messagebox
from openpyxl import Workbook
from openpyxl import load_workbook
import os

filename = "student_scores.xlsx"

# Create Excel file 
def excel_file():
    if not os.path.exists(filename):
        work_book = Workbook()
        work_sheet = work_book.active
        work_sheet.title = "StudentScore"
        work_sheet.append(["Name", "Score", "Result"])
        work_book.save(filename)

def save_record():
    name = name_entry.get()
    try:
        score = float(score_entry.get())
    except ValueError:
        messagebox.showerror("Invalid input", "Please enter a valid number for the score.")
        return

    if 1 <= score <= 74:
        result = "Fail"
    elif 75 <= score <= 100:
        result = "Pass"
    else:
        messagebox.showerror("Invalid input", "Score must be between 1 and 100.")
        return

    work_book = load_workbook(filename)
    work_sheet = work_book.active
    next_row = work_sheet.max_row + 1
    work_sheet[f"A{next_row}"] = name
    work_sheet[f"B{next_row}"] = score
    work_sheet[f"C{next_row}"] = result
    work_book.save(filename)

    messagebox.showinfo("Success", "Record saved!")
    name_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)

# Initialize Excel file
excel_file()

# TKINTER UI
window = tk.Tk()
window.title("Student Score Tracker")
window.geometry("500x300")

frame = tk.Frame(window)
frame.pack(expand=True)

tk.Label(frame, text="Name").grid(row=0, column=0, pady=5, sticky="e")
name_entry = tk.Entry(frame)
name_entry.grid(row=0, column=1, padx=10, pady=5, sticky="w")

tk.Label(frame, text="Score").grid(row=1, column=0, pady=5, sticky="e")
score_entry = tk.Entry(frame)
score_entry.grid(row=1, column=1, padx=10, pady=5, sticky="w")

tk.Button(frame, text="Save Record", command=save_record).grid(row=3, column=0, columnspan=2, pady=20)

window.mainloop()


    


